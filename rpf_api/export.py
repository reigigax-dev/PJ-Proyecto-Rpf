import logging
from datetime import date, datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

SEX_LABEL = {True: "F", False: "M"}

COLUMNS = [
    # (título, ancho)
    ("ID", 7),
    ("RUT", 14),
    ("Nombres", 26),
    ("Apellidos", 26),
    ("Sexo", 7),
    ("Fecha de nacimiento", 20),
    ("Lugar de nacimiento", 22),
    ("N° documento", 15),
    ("Foto frontal", 16),
    ("Foto trasera", 16),
]
COL_DATE, COL_DOC_NUMBER, COL_FRONT, COL_BACK = 6, 8, 9, 10


def format_rut(rut):
    clean = "".join(c for c in str(rut or "") if c.isalnum()).upper()
    if len(clean) < 2 or not clean[:-1].isdigit():
        return clean
    return f"{int(clean[:-1]):,}".replace(",", ".") + f"-{clean[-1]}"


def _text(value):
    if value is None:
        return ""
    return ILLEGAL_CHARACTERS_RE.sub("", str(value)).strip()


def _number(value):
    if value in (None, ""):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return _text(value) or None


def _date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value:
        try:
            return date.fromisoformat(str(value)[:10])
        except ValueError:
            return _text(value)
    return None


def _image_url(field, image_url):
    if image_url is None or not field or not getattr(field, "name", None):
        return None
    try:
        return image_url(field)
    except Exception:  # archivo sin nombre, storage mal configurado, etc.
        logger.warning("No se pudo obtener la URL de %r", getattr(field, "name", field))
        return None


def _row_values(r):
    return [
        r.id,
        format_rut(getattr(r, "document", "")),
        _text(getattr(r, "first_names", "")),
        _text(getattr(r, "last_names", "")),
        SEX_LABEL.get(getattr(r, "sex", None), ""),
        _date(getattr(r, "born_date", None)),
        _text(getattr(r, "born_location", None)),
        _number(getattr(r, "document_number", None)),
    ]


def build_workbook(records, image_url=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Registros"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for col, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width

    row = 1
    for r in records:
        row += 1
        try:
            values = _row_values(r)
        except Exception:
            logger.exception("Registro %s con datos inválidos para el Excel", getattr(r, "id", "?"))
            values = [getattr(r, "id", None), "(error al leer este registro)"]

        for col, value in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=value)
        ws.cell(row=row, column=COL_DATE).number_format = "DD-MM-YYYY"
        ws.cell(row=row, column=COL_DOC_NUMBER).number_format = "0"

        for col, attr in ((COL_FRONT, "document_front_picture"), (COL_BACK, "document_back_picture")):
            url = _image_url(getattr(r, attr, None), image_url)
            if url:
                cell = ws.cell(row=row, column=col, value="Ver foto")
                cell.hyperlink = url
                cell.font = Font(color="0563C1", underline="single")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(row, 1)}"

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
